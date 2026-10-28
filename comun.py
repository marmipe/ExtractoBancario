"""Shared helpers: read and validate a BCP or Interbank export, write the DTW trama, query SQL Server."""
import re
import subprocess
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

AQUI = Path(__file__).parent
CREDENCIALES = AQUI.parents[1] / "mssql_credentials.env"  # SQL-MCP-SX\mssql_credentials.env

# Bank accounts of Shouxin -> (SAP G/L account, currency). This table is the only source:
# the program does not look the account up in SAP. Taken from the company's bank accounts
# (house banks, agreed with the SAP analyst, 2026-10-27). Transit, time-deposit and BVLETRAS
# entries are left out on purpose. Matching ignores dashes, dots and spaces, so
# "100-301150000818" and "100301150000818" are the same account.
CUENTAS = {
    "191-1913747-0-76": ("104111", "S/"),     # BCP M.N.
    "191-1870184-1-56": ("104112", "US$"),    # BCP M.E.
    "377-3005555475": ("104123", "S/"),       # Interbank M.N.
    "377-3005555482": ("104124", "US$"),      # Interbank M.E.
    "00-002-107104": ("104211", "S/"),        # Banco de la Nación M.N. (detracciones)
    "868059": ("104115", "S/"),               # ICBC M.N.
    "889152": ("104116", "S/"),               # ICBC M.N.
    "06000-101000000889152": ("104116", "S/"),
    "852654": ("104117", "US$"),              # ICBC M.E.
    "06000-101000000852654": ("104117", "US$"),
    "873757": ("104118", "US$"),              # ICBC M.E.
    "894850": ("104119", "US$"),              # ICBC M.E.
    "100-301150000818": ("104121", "S/"),     # Bank of China M.N.
    "100-301150000807": ("104122", "US$"),    # Bank of China M.E.
    "100-301150002861": ("104125", "CNY"),    # Bank of China yuanes
}

# Currency as the bank prints it in B2 (lowercase, no accents, start of the text) -> currency code.
MONEDAS = {"soles": "S/", "dolares": "US$", "yuan": "CNY"}


def _solo_digitos(texto: str) -> str:
    return "".join(c for c in texto if c.isdigit())


_CUENTAS_NORM = {_solo_digitos(k): v for k, v in CUENTAS.items()}
assert len(_CUENTAS_NORM) == len(CUENTAS), "CUENTAS tiene dos claves con los mismos dígitos"


def _sin_tildes(texto) -> str:
    return unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().strip().lower()


def _moneda_sap(texto: str) -> str | None:
    """Find a currency word anywhere in the text ('Soles', 'Corriente Dólares 377-...')."""
    for palabra in _sin_tildes(texto).split():
        for nombre, sap in MONEDAS.items():
            if palabra.startswith(nombre):
                return sap
    return None


HEADER_ROW = 5  # BCP: 'Fecha', 'Fecha valuta', ...; data starts on the next row
TRAMA_HEADER = (
    "Sequence\tAccountCode\tReference\tDueDate\tMemo\tDebitAmount\tCreditAmount\n"
    "Sequence\tAccountCode\tNUMERO\tDueDate\tDESCRIPCION\tDebitAmount\tCreditAmount\n"
)


class ErrorExtracto(Exception):
    pass


def _cuenta(cuenta_banco: str, texto_moneda: str, celda_moneda: str) -> tuple[str, str]:
    """Bank account + currency text from the file -> (SAP account, currency), or refuse."""
    if _solo_digitos(cuenta_banco) not in _CUENTAS_NORM:
        raise ErrorExtracto(f"Cuenta bancaria {cuenta_banco!r} no está configurada en CUENTAS (comun.py).")
    cuenta_sap, moneda = _CUENTAS_NORM[_solo_digitos(cuenta_banco)]
    moneda_archivo = _moneda_sap(texto_moneda)
    if moneda_archivo is None:
        raise ErrorExtracto(f"Moneda {texto_moneda!r} ({celda_moneda}) no reconocida. "
                            f"Agréguela en MONEDAS (comun.py).")
    if moneda_archivo != moneda:
        raise ErrorExtracto(f"La cuenta {cuenta_banco} es {moneda} ({cuenta_sap}), "
                            f"pero el archivo dice moneda {texto_moneda!r} ({celda_moneda}).")
    return cuenta_sap, moneda


def _numero(valor, fila: int, campo: str) -> float | None:
    if valor in (None, ""):
        return None
    try:
        return round(float(valor), 2)
    except (TypeError, ValueError):
        raise ErrorExtracto(f"Fila {fila}: {campo} inválido ({valor!r}).")


def _fecha(valor, fila: int) -> datetime:
    if isinstance(valor, datetime):
        return valor
    try:
        return datetime.strptime(str(valor).strip(), "%d/%m/%Y")
    except ValueError:
        raise ErrorExtracto(f"Fila {fila}: fecha inválida ({valor!r}).")


def _leer_bcp(ws) -> dict:
    """BCP 'movimientos históricos': account in B1, currency in B2, header in row 5, signed Monto + Saldo."""
    cuenta_banco = str(ws.cell(1, 2).value or "").split(" - ")[0].strip()
    cuenta_sap, moneda = _cuenta(cuenta_banco, str(ws.cell(2, 2).value or ""), "B2")

    encabezado = [c.value for c in ws[HEADER_ROW]]
    try:
        col = {n: encabezado.index(n) for n in
               ("Fecha", "Descripción operación", "Monto", "Saldo", "Operación - Número", "UTC")}
    except ValueError as e:
        raise ErrorExtracto(f"Fila {HEADER_ROW} no tiene el encabezado esperado de BCP: {e}")

    crudas = []
    for n, fila in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True), HEADER_ROW + 1):
        if not fila[col["Fecha"]]:
            continue
        monto = _numero(fila[col["Monto"]], n, "monto")
        saldo = _numero(fila[col["Saldo"]], n, "saldo")
        if monto is None or saldo is None:
            raise ErrorExtracto(f"Fila {n}: monto o saldo vacío.")
        numero = str(fila[col["Operación - Número"]] or "").strip()
        ref = numero if numero.strip("0") else str(fila[col["UTC"]] or "").strip()
        crudas.append(dict(fila_excel=n, fecha=_fecha(fila[col["Fecha"]], n), monto=monto, saldo=saldo,
                           ref=ref, memo=str(fila[col["Descripción operación"]] or "").strip()))
    return dict(banco="BCP", cuenta_banco=cuenta_banco, cuenta_sap=cuenta_sap, moneda=moneda, crudas=crudas)


def _leer_interbank(ws) -> dict:
    """Interbank 'Consulta de Movimientos de Cuenta Corriente'.

    'Cuenta:' label followed by e.g. 'Corriente Dólares 377-3005555482'; header row starting
    'Fecha de operación'; Cargo (negative) and Abono (positive) in separate columns. 'Saldo
    contable' is usually empty, so balances are only checked when the bank fills it.
    """
    etiquetas, fila_enc = {}, None
    for fila in ws.iter_rows(min_row=1, max_row=30):
        celdas = [c for c in fila if c.value not in (None, "")]
        for i, c in enumerate(celdas):
            t = _sin_tildes(c.value)
            if t in ("cuenta:", "rango de fechas:") and i + 1 < len(celdas):
                etiquetas[t] = celdas[i + 1]
            if t == "fecha de operacion" and fila_enc is None:
                fila_enc = c.row
    if "cuenta:" not in etiquetas or fila_enc is None:
        raise ErrorExtracto("No se encontró 'Cuenta:' o el encabezado 'Fecha de operación' de Interbank.")
    celda_cuenta = etiquetas["cuenta:"]
    texto_cuenta = str(celda_cuenta.value).strip()
    cuenta_banco = texto_cuenta.split()[-1]
    cuenta_sap, moneda = _cuenta(cuenta_banco, texto_cuenta, celda_cuenta.coordinate)

    encabezado = [_sin_tildes(c.value) for c in ws[fila_enc]]
    nombres = {"fecha": "fecha de operacion", "numero": "nro. de operacion", "mov": "movimiento",
               "desc": "descripcion", "cargo": "cargo", "abono": "abono", "saldo": "saldo contable"}
    try:
        col = {k: encabezado.index(v) for k, v in nombres.items()}
    except ValueError as e:
        raise ErrorExtracto(f"Fila {fila_enc} no tiene el encabezado esperado de Interbank: {e}")

    crudas = []
    for n, fila in enumerate(ws.iter_rows(min_row=fila_enc + 1, values_only=True), fila_enc + 1):
        if not fila[col["fecha"]]:
            continue
        cargo = _numero(fila[col["cargo"]], n, "cargo")
        abono = _numero(fila[col["abono"]], n, "abono")
        if (cargo is None) == (abono is None):
            raise ErrorExtracto(f"Fila {n}: debe tener cargo o abono (solo uno).")
        if (cargo is not None and cargo > 0) or (abono is not None and abono < 0):
            raise ErrorExtracto(f"Fila {n}: signo inesperado (cargo debe ser negativo, abono positivo).")
        numero = str(fila[col["numero"]] or "").strip()
        mov = str(fila[col["mov"]] or "").strip()
        desc = str(fila[col["desc"]] or "").strip()
        if not numero.strip("-0 "):
            # ITF and similar lines have no number of their own; they name the operation they belong to.
            m = re.search(r"OPERACION:\s*(\d+)", desc, re.IGNORECASE)
            numero = m.group(1) if m else ""
        # Memo as treasury types it: the Descripción, except ITF lines ('OPERACION: n') which say 'ITF'.
        if _sin_tildes(mov) == "itf":
            memo = "ITF"
        else:
            memo = desc or mov
        crudas.append(dict(fila_excel=n, fecha=_fecha(fila[col["fecha"]], n),
                           monto=cargo if cargo is not None else abono,
                           saldo=_numero(fila[col["saldo"]], n, "saldo"),
                           ref=numero, memo=memo))

    # The file states the range it covers ('Desde: dd/mm/yyyy', next row 'Hasta: dd/mm/yyyy').
    if "rango de fechas:" in etiquetas:
        celda = etiquetas["rango de fechas:"]
        siguiente = " ".join(str(c.value) for c in ws[celda.row + 1] if c.value)
        desde = re.search(r"(\d{2}/\d{2}/\d{4})", str(celda.value))
        hasta = re.search(r"Hasta:\s*(\d{2}/\d{2}/\d{4})", siguiente)
        if desde and hasta:
            d, h = (datetime.strptime(x.group(1), "%d/%m/%Y") for x in (desde, hasta))
            for c in crudas:
                if not d <= c["fecha"] <= h:
                    raise ErrorExtracto(f"Fila {c['fila_excel']}: fecha {c['fecha']:%d/%m/%Y} fuera del "
                                        f"rango {d:%d/%m/%Y}–{h:%d/%m/%Y} que indica el archivo.")
    return dict(banco="Interbank", cuenta_banco=cuenta_banco, cuenta_sap=cuenta_sap, moneda=moneda,
                crudas=crudas)


def leer_extracto(origen: Path) -> dict:
    """Read a BCP or Interbank export (detected from its layout) and validate it.

    Returns banco, cuenta_banco, cuenta_sap, moneda, cuadra (True when running balances were
    checked), saldo_inicial / saldo_final (None when the file has no balances) and lineas
    (oldest first; each a dict with fecha 'YYYYMMDD', ref, memo, debe, haber).
    """
    ws = openpyxl.load_workbook(origen, data_only=True).worksheets[0]
    if _sin_tildes(ws.cell(1, 1).value) == "cuenta" and _sin_tildes(ws.cell(2, 1).value) == "moneda":
        ext = _leer_bcp(ws)
    elif any("consulta de movimientos" in _sin_tildes(c.value)
             for fila in ws.iter_rows(max_row=10) for c in fila if c.value):
        ext = _leer_interbank(ws)
    else:
        raise ErrorExtracto("Formato no reconocido: no es una exportación de BCP ni de Interbank.")

    crudas = ext.pop("crudas")  # bank order: newest first
    if not crudas:
        raise ErrorExtracto("El archivo no tiene movimientos.")
    for c in crudas:
        if not c["memo"]:
            raise ErrorExtracto(f"Fila {c['fila_excel']}: descripción vacía.")
        if c["monto"] == 0:
            raise ErrorExtracto(f"Fila {c['fila_excel']}: monto cero.")
    for nuevo, viejo in zip(crudas, crudas[1:]):
        if nuevo["fecha"] < viejo["fecha"]:
            raise ErrorExtracto(f"Fila {nuevo['fila_excel']}: fechas fuera de orden.")

    # Each balance must equal the next (older) balance plus this movement. Checked whenever every
    # line has a balance (always for BCP; Interbank usually leaves 'Saldo contable' empty).
    cuadra = all(c["saldo"] is not None for c in crudas)
    if cuadra:
        for nuevo, viejo in zip(crudas, crudas[1:]):
            if round(viejo["saldo"] + nuevo["monto"], 2) != nuevo["saldo"]:
                raise ErrorExtracto(
                    f"Fila {nuevo['fila_excel']}: el saldo no cuadra "
                    f"({viejo['saldo']:,.2f} + {nuevo['monto']:,.2f} != {nuevo['saldo']:,.2f}).")

    ext["lineas"] = [dict(fecha=c["fecha"].strftime("%Y%m%d"), ref=c["ref"][:27], memo=c["memo"][:254],
                          debe=-c["monto"] if c["monto"] < 0 else 0.0,
                          haber=c["monto"] if c["monto"] > 0 else 0.0)
                     for c in reversed(crudas)]
    ext["cuadra"] = cuadra
    ext["saldo_inicial"] = round(crudas[-1]["saldo"] - crudas[-1]["monto"], 2) if cuadra else None
    ext["saldo_final"] = crudas[0]["saldo"] if cuadra else None
    return ext


leer_bcp = leer_extracto  # old name, still used by crear_prueba.py


def resumen(ext: dict) -> str:
    """Two-line summary of a validated statement for the console and the log."""
    total_debe = sum(l["debe"] for l in ext["lineas"])
    total_haber = sum(l["haber"] for l in ext["lineas"])
    texto = (f"{ext['banco']} {ext['cuenta_banco']} -> {ext['cuenta_sap']} ({ext['moneda']}), "
             f"{len(ext['lineas'])} líneas, {ext['lineas'][0]['fecha']} a {ext['lineas'][-1]['fecha']}\n"
             f"Cargos {total_debe:,.2f}  Abonos {total_haber:,.2f}  ")
    if ext["cuadra"]:
        return texto + f"Saldo {ext['saldo_inicial']:,.2f} -> {ext['saldo_final']:,.2f}  (cuadra)"
    return texto + "(el archivo no trae saldo por línea: no se pudo cuadrar contra el banco)"


def escribir_trama(extracto: dict, destino: Path) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    with open(destino, "w", encoding="utf-8", newline="\r\n") as f:
        f.write(TRAMA_HEADER)
        for l in extracto["lineas"]:
            f.write("\t".join(["", extracto["cuenta_sap"], l["ref"], l["fecha"], l["memo"],
                               f"{l['debe']:.2f}", f"{l['haber']:.2f}"]) + "\n")


def sql(base: str, consulta: str) -> list[list[str]]:
    """Run a read-only query with sqlcmd; returns rows as lists of strings."""
    cfg = {}
    for linea in CREDENCIALES.read_text(encoding="utf-8").splitlines():
        linea = linea.strip()
        if linea and not linea.startswith("#") and "=" in linea:
            k, v = linea.split("=", 1)
            cfg[k] = v
    r = subprocess.run(
        ["sqlcmd", "-S", f"{cfg['MSSQL_SERVER']},{cfg.get('MSSQL_PORT', '1433')}",
         "-U", cfg["MSSQL_USER"], "-P", cfg["MSSQL_PASSWORD"], "-d", base,
         "-W", "-h", "-1", "-s", "\t", "-b", "-Q", "SET NOCOUNT ON; " + consulta],
        capture_output=True, text=True, encoding="cp1252")
    if r.returncode != 0:
        raise RuntimeError(f"sqlcmd falló: {r.stdout}{r.stderr}")
    return [l.split("\t") for l in r.stdout.splitlines() if l.strip()]
